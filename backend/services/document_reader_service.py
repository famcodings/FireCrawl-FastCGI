"""Document reader service for extracting text from various file formats."""
import os
from typing import List, Dict, Tuple
from io import BytesIO
import PyPDF2
import docx
import openpyxl
from config import config


class DocumentReaderService:
    """Service for reading and extracting text from various document formats."""
    
    def extract_text_from_pdf(self, file_content: bytes) -> str:
        """Extract text from PDF file."""
        try:
            pdf_reader = PyPDF2.PdfReader(BytesIO(file_content))
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            return text.strip()
        except Exception as e:
            raise ValueError(f"Error extracting text from PDF: {str(e)}")
    
    def extract_text_from_docx(self, file_content: bytes) -> str:
        """Extract text from DOCX file."""
        try:
            doc = docx.Document(BytesIO(file_content))
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text.strip()
        except Exception as e:
            raise ValueError(f"Error extracting text from DOCX: {str(e)}")
    
    def extract_text_from_txt(self, file_content: bytes) -> str:
        """Extract text from TXT file."""
        try:
            return file_content.decode('utf-8').strip()
        except UnicodeDecodeError:
            try:
                return file_content.decode('latin-1').strip()
            except Exception as e:
                raise ValueError(f"Error extracting text from TXT: {str(e)}")
    
    def extract_text_from_excel(self, file_content: bytes) -> str:
        """Extract text from Excel file."""
        try:
            workbook = openpyxl.load_workbook(BytesIO(file_content))
            text = ""
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text += f"Sheet: {sheet_name}\n"
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join(str(cell) for cell in row if cell is not None)
                    if row_text.strip():
                        text += row_text + "\n"
            return text.strip()
        except Exception as e:
            raise ValueError(f"Error extracting text from Excel: {str(e)}")
    
    def extract_text_from_file(self, file_content: bytes, file_extension: str) -> str:
        """Extract text from uploaded file based on file extension."""
        file_extension = file_extension.lower()
        
        if file_extension == '.pdf':
            return self.extract_text_from_pdf(file_content)
        elif file_extension == '.docx':
            return self.extract_text_from_docx(file_content)
        elif file_extension == '.txt':
            return self.extract_text_from_txt(file_content)
        elif file_extension in ['.xlsx', '.xls']:
            return self.extract_text_from_excel(file_content)
        else:
            raise ValueError(f"Unsupported file type: {file_extension}")
    
    def process_files(self, files: List[Tuple[bytes, str]]) -> Dict:
        """Process multiple uploaded files and extract text content."""
        try:
            documents = []
            processed_files = []
            
            for file_content, filename in files:
                file_extension = os.path.splitext(filename)[1]
                
                if file_extension not in config.ALLOWED_FILE_TYPES:
                    return {
                        "success": False,
                        "error": f"Unsupported file type: {file_extension}. Allowed types: {', '.join(config.ALLOWED_FILE_TYPES)}"
                    }
                
                if len(file_content) > config.MAX_FILE_SIZE:
                    return {
                        "success": False,
                        "error": f"File {filename} is too large. Maximum size: {config.MAX_FILE_SIZE / (1024*1024):.1f}MB"
                    }
                
                try:
                    # Extract text from the file
                    text = self.extract_text_from_file(file_content, file_extension)
                    if text.strip():
                        # Create document object
                        documents.append({
                            "text": text,
                            "metadata": {
                                "filename": filename,
                                "type": file_extension
                            }
                        })
                        processed_files.append(filename)
                except Exception as e:
                    return {
                        "success": False,
                        "error": f"Error processing file {filename}: {str(e)}"
                    }
            
            if not documents:
                return {
                    "success": False,
                    "error": "No readable text found in any of the uploaded documents"
                }
            
            return {
                "success": True,
                "documents": documents,
                "processed_files": processed_files
            }
            
        except Exception as e:
            return {
                "success": False,
                "error": f"Error processing documents: {str(e)}"
            }


# Global document reader service instance
document_reader_service = DocumentReaderService()
